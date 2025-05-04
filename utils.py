import random
import difflib
import re
import pytz
from langchain.chat_models import ChatOpenAI
from PyPDF2 import PdfReader
from docx import Document
from fpdf import FPDF
from datetime import datetime
from io import BytesIO

# Greeting and farewell keyword libraries
GREETING_KEYWORDS = {"hi", "hello", "hey", "how are you", "good morning", "good afternoon", "greetings"}
FAREWELL_KEYWORDS = {"bye", "goodbye", "thank you", "thanks", "see you", "take care", "farewell"}

# Response libraries
GREETING_REPLIES = [
    "Hello! How can I help you today?",
    "Hi there! What would you like to know?",
    "Hey! I'm here to help you with your PDFs.",
    "Greetings! Ready when you are.",
    "Hi! Ask me anything from your documents."
]

FAREWELL_REPLIES = [
    "Goodbye! Let me know if you need help again.",
    "Take care! Come back anytime.",
    "You're welcome! Have a great day!",
    "Bye! It was nice chatting with you.",
    "Thanks! I'm here if you need me again."
]

def handle_greeting(user_input: str):
    lower_input = user_input.lower().strip()
    if any(keyword in lower_input for keyword in GREETING_KEYWORDS) and len(lower_input.split()) <= 4:
        return random.choice(GREETING_REPLIES)
    return None

def handle_farewell(user_input: str):
    lower_input = user_input.lower().strip()
    if any(keyword in lower_input for keyword in FAREWELL_KEYWORDS) and len(lower_input.split()) <= 4:
        return random.choice(FAREWELL_REPLIES)
    return None

def get_labeled_documents_from_any(uploaded_docs):
    labeled_docs = []
    for i, doc in enumerate(uploaded_docs):
        name = doc.name.lower()
        text = ""

        if name.endswith(".pdf"):
            pdf_reader = PdfReader(doc)
            for page in pdf_reader.pages:
                content = page.extract_text()
                if content:
                    text += content

        elif name.endswith(".docx"):
            word_doc = Document(doc)
            for para in word_doc.paragraphs:
                text += para.text + "\n"

        elif name.endswith(".txt"):
            text += doc.read().decode("utf-8") + "\n"

        else:
            continue

        label = f"Document {i+1}: {doc.name}"
        labeled_docs.append({"label": label, "text": text})
    
    return labeled_docs

def summarize_documents(labeled_docs):
    llm = ChatOpenAI(model_name="gpt-3.5-turbo", temperature=0.3)

    summaries = []
    for doc in labeled_docs:
        prompt = f"Please summarize the following document:\n\n{doc['text'][:4000]}\n\nSummary:"
        response = llm.predict(prompt)
        summaries.append({"label": doc["label"], "summary": response.strip()})
    return summaries

def is_summary_question(question: str) -> bool:
    summary_keywords = ["summary", "summarize", "overview", "main idea", "key points", "what is in", "content of"]
    return any(keyword in question.lower() for keyword in summary_keywords)

def extract_target_doc_label(question: str, docs: list, cutoff: float = 0.6) -> str:
    question_lower = question.lower()
    labels = [item['label'].lower() for item in docs]

    doc_num_match = re.search(r"(document|doc|file)\s*(\d+)", question_lower)
    if doc_num_match:
        doc_num = int(doc_num_match.group(2))
        if 1 <= doc_num <= len(docs):
            return docs[doc_num - 1]["label"]
    
    matches = difflib.get_close_matches(question_lower, labels, n=1, cutoff=cutoff)

    if matches:
        matched_label_lower = matches[0]
        for item in docs:
            if item['label'].lower() == matched_label_lower:
                return item['label']
    
    for item in docs:
        if item['label'].lower() in question_lower:
            return item['label']
    
    return None

def is_wordcount_question(question: str) -> bool:
    keywords = ["word count", "how many words", "number of words", "words in", "wordcount", "word"]
    return any(keyword in question.lower() for keyword in keywords)

def count_words_in_documents(labeled_docs):
    word_counts = []
    for doc in labeled_docs:
        word_count = len(doc["text"].split())
        word_counts.append({
            "label": doc["label"],
            "word_count": word_count
        })
    return word_counts

def estimate_multicell_height(pdf, text, width, line_height):
    lines = pdf.multi_cell(width, line_height, text, split_only=True)
    return len(lines) * line_height + 4 

def save_chat_to_pdf(chat_history):
    def strip_emojis(text):
        return re.sub(r'[^\x00-\x7F]+', '', text)

    def remove_newlines(text):
        return re.sub(r'\s*\n\s*', ' ', text.strip())

    pdf = FPDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=False)
    page_height = 297  # A4 height in mm
    margin_top = 10
    margin_bottom = 10
    usable_height = page_height - margin_top - margin_bottom
    line_height = 8
    box_spacing = 6
    box_width = 190

    # Header
    pdf.set_font("Arial", 'B', 16)
    pdf.cell(0, 10, "Chat History", ln=True, align="C")
    pdf.set_font("Arial", '', 10)
    malaysia_time = datetime.now(pytz.timezone("Asia/Kuala_Lumpur")).strftime("%B %d, %Y %H:%M")
    pdf.cell(0, 10, f"Exported on {malaysia_time} (MYT)", ln=True, align="C")
    pdf.ln(5)
    pdf.set_font("Arial", '', 12)

    # Pair user and assistant messages
    i = 0
    while i < len(chat_history):
        entry = chat_history[i]
        if entry["role"] == "user":
            user_msg = strip_emojis(entry["content"]).strip()
            label_user = f"You:\n{user_msg}"

            # Look ahead to the assistant's reply
            assistant_msg = ""
            if i + 1 < len(chat_history) and chat_history[i + 1]["role"] == "assistant":
                assistant_msg = remove_newlines(strip_emojis(chat_history[i + 1]["content"]).strip())
                i += 1  # Skip assistant entry on next loop
            # assistant_msg = re.sub(r"[|\#\-\t]", " ", assistant_msg)
            # assistant_msg = re.sub(r"\s+", " ", assistant_msg).strip()  # Normalize spaces
            assistant_msg = re.sub(r"[|\#\-\t]", " ", assistant_msg)
            assistant_msg = re.sub(r"[ ]{2,}", " ", assistant_msg)

            label_assistant = f"Assistant:\n{assistant_msg}"

            # Estimate box heights
            user_box_height = estimate_multicell_height(pdf, label_user, box_width, line_height)
            assistant_box_height = estimate_multicell_height(pdf, label_assistant, box_width, line_height)
            total_pair_height = user_box_height + assistant_box_height + box_spacing

            # Add new page if not enough space
            if pdf.get_y() + total_pair_height > usable_height:
                pdf.add_page()

            # Render user message
            y_start = pdf.get_y()
            pdf.rect(10, y_start, box_width, user_box_height)
            pdf.set_xy(12, y_start + 2)
            pdf.multi_cell(0, line_height, label_user)
            pdf.ln(2)

            # Render assistant message
            y_start = pdf.get_y()
            pdf.rect(10, y_start, box_width, assistant_box_height)
            pdf.set_xy(12, y_start + 2)
            pdf.set_text_color(0, 102, 204)
            pdf.multi_cell(0, line_height, label_assistant)
            pdf.set_text_color(0, 0, 0)
            pdf.ln(4)

        i += 1

    # Output PDF
    pdf_bytes = pdf.output(dest='S').encode('latin1')
    return BytesIO(pdf_bytes)
